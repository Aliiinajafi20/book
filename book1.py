from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = Workbook()
ws = wb.active

data = [
    ('ali', 'najafi', 9336139966),
    ['shirin', 'alizade', 256555443000],
    ['reza', 'bahadori', 9498496000],
    ['yasin', 'nobakht', 34498494700],
    ['javad', 'mosavi', '8788494700'],
    ['hadi', 'vaezi', 42418494700],
    ['mohsen', 'heydari', 585896564700],
    ['mohammad', 'rostami', 6998494700],
    ['sara', 'khezri', 3449887874700],
    ['maryam', 'jalali', 7788494700],
    ['morteza', 'bahrami', 34498444880],
    ['kamran', 'manafi', 4878494700],
]

# add column headings. NB. these must be strings
ws.append(["name", "lastname","phone"])
for row in data:
    ws.append(row)

tab = Table(displayName="Table1", ref="A1:C13")

# Add a default style with striped rows and banded columns
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=True)
tab.tableStyleInfo = style

'''
Table must be added using ws.add_table() method to avoid duplicate names.
Using this method ensures table name is unque through out defined names and all other table name. 
'''
ws.add_table(tab)
wb.save("table.xlsx")